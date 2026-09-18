from __future__ import annotations

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from witness_engine.ingestion import (
    extract_document,
    supported_extensions,
)
from witness_engine.pipeline import index_document
from witness_engine.retrieval import LocalEvidenceIndex


def test_v1_registry_includes_all_published_file_targets():
    extensions = set(supported_extensions())
    for suffix in (
        ".pdf",
        ".md",
        ".txt",
        ".html",
        ".htm",
        ".docx",
        ".pptx",
        ".xlsx",
        ".csv",
    ):
        assert suffix in extensions


def test_html_extracts_structural_blocks_and_table_rows(tmp_path):
    source = tmp_path / "saved-page.html"
    source.write_text(
        """
        <html><head><title>Ignored title</title>
        <style>.hidden{display:none}</style></head>
        <body>
          <h1>Deployment</h1>
          <p>The current API port is 5200.</p>
          <table>
            <tr><th>Service</th><th>Port</th></tr>
            <tr><td>API</td><td>5200</td></tr>
          </table>
          <script>ATTACK_SHOULD_NOT_BE_EVIDENCE</script>
        </body></html>
        """,
        encoding="utf-8",
    )

    document = extract_document(source, "source-html")
    assert document.media_type == "text/html"
    by_locator = {block.locator: block for block in document.blocks}
    assert by_locator["element:h1:1"].text == "Deployment"
    assert by_locator["element:p:1"].text == "The current API port is 5200."
    assert by_locator["table:1/row:1"].text == "Service | Port"
    assert by_locator["table:1/row:2"].text == "API | 5200"
    assert all(
        "ATTACK_SHOULD_NOT_BE_EVIDENCE" not in block.text
        for block in document.blocks
    )

    with LocalEvidenceIndex(tmp_path / "html.sqlite3") as index:
        result = index_document(source, index)
        hits = index.search("current API port")
    assert result.media_type == "text/html"
    assert hits
    assert hits[0].locator == "element:p:1"


def test_csv_rows_are_searchable_with_stable_row_locators(tmp_path):
    source = tmp_path / "services.csv"
    source.write_text(
        "service,port\napi,5200\nauth,7000\n",
        encoding="utf-8",
    )
    document = extract_document(source, "source-csv")
    assert document.media_type == "text/csv"
    assert [block.locator for block in document.blocks] == [
        "row:1",
        "row:2",
        "row:3",
    ]
    assert document.blocks[1].text == "api | 5200"

    with LocalEvidenceIndex(tmp_path / "csv.sqlite3") as index:
        index_document(source, index)
        hits = index.search("api 5200")
    assert hits
    assert hits[0].locator == "row:2"


def test_xlsx_rows_are_searchable_with_sheet_row_locators(tmp_path):
    source = tmp_path / "services.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Services"
    sheet.append(["service", "port"])
    sheet.append(["api", 5200])
    sheet.append(["auth", 7000])
    workbook.save(source)

    document = extract_document(source, "source-xlsx")
    assert document.media_type.endswith(
        "spreadsheetml.sheet"
    )
    assert [block.locator for block in document.blocks] == [
        "sheet:1/row:1",
        "sheet:1/row:2",
        "sheet:1/row:3",
    ]
    assert document.blocks[1].text == "api | 5200"

    with LocalEvidenceIndex(tmp_path / "xlsx.sqlite3") as index:
        index_document(source, index)
        hits = index.search("auth 7000")
    assert hits
    assert hits[0].locator == "sheet:1/row:3"


def test_pptx_text_and_table_rows_are_searchable_with_slide_locators(
    tmp_path,
):
    source = tmp_path / "architecture.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )
    box = slide.shapes.add_textbox(
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(1),
    )
    box.text = "The API gateway listens on port 5200."
    table_shape = slide.shapes.add_table(
        2,
        2,
        Inches(1),
        Inches(2),
        Inches(5),
        Inches(1.5),
    )
    table = table_shape.table
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "API"
    table.cell(1, 1).text = "Platform"
    presentation.save(source)

    document = extract_document(source, "source-pptx")
    assert document.media_type.endswith(
        "presentationml.presentation"
    )
    by_locator = {block.locator: block for block in document.blocks}
    assert (
        by_locator["slide:1/shape:1"].text
        == "The API gateway listens on port 5200."
    )
    assert by_locator["slide:1/table:1/row:1"].text == "Service | Owner"
    assert by_locator["slide:1/table:1/row:2"].text == "API | Platform"

    with LocalEvidenceIndex(tmp_path / "pptx.sqlite3") as index:
        index_document(source, index)
        hits = index.search("gateway port 5200")
    assert hits
    assert hits[0].locator == "slide:1/shape:1"
