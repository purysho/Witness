from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from witness_engine.ingestion import UnsupportedDocumentError, extract_document
from witness_engine.pipeline import index_document, search_evidence
from witness_engine.retrieval import LocalEvidenceIndex


def test_markdown_preserves_heading_and_line_provenance(tmp_path: Path):
    source = tmp_path / "notes.md"
    source.write_text(
        "# Authentication\n\nBearer tokens protect the API.\n\n## Rotation\n\nRotate keys weekly.\n",
        encoding="utf-8",
    )

    with LocalEvidenceIndex(tmp_path / "witness.sqlite3") as index:
        result = index_document(source, index)
        hits = search_evidence("Bearer tokens", index)

    assert result.media_type == "text/markdown"
    assert result.block_count == 4
    assert hits
    assert hits[0].source_version_id == result.source_version_id
    assert hits[0].locator == "line:3"


def test_source_code_is_searchable_with_line_range_locator(tmp_path: Path):
    source = tmp_path / "auth.py"
    source.write_text(
        "def issue_token(user_id: str):\n"
        "    token = sign(user_id)\n"
        "    return token\n",
        encoding="utf-8",
    )

    with LocalEvidenceIndex(tmp_path / "witness.sqlite3") as index:
        result = index_document(source, index)
        hits = search_evidence("issue_token", index)

    assert result.media_type == "text/x-source-code"
    assert hits
    assert hits[0].source_version_id == result.source_version_id
    assert hits[0].locator == "line:1-3"


def test_docx_paragraphs_and_table_rows_are_searchable(tmp_path: Path):
    source = tmp_path / "design.docx"
    document = Document()
    document.add_heading("Authentication", level=1)
    document.add_paragraph("The gateway validates bearer tokens.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Rotation"
    table.rows[0].cells[1].text = "Weekly"
    document.save(source)

    with LocalEvidenceIndex(tmp_path / "witness.sqlite3") as index:
        result = index_document(source, index)
        paragraph_hits = search_evidence("gateway bearer", index)
        table_hits = search_evidence("Rotation Weekly", index)

    assert result.media_type.startswith("application/vnd.openxmlformats")
    assert paragraph_hits
    assert paragraph_hits[0].locator == "paragraph:2"
    assert table_hits
    assert table_hits[0].locator == "table:1/row:1"


def test_pdf_pages_without_text_report_warning_not_fake_evidence(tmp_path: Path):
    source = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with source.open("wb") as handle:
        writer.write(handle)

    document = extract_document(source, "source-version-test")

    assert document.media_type == "application/pdf"
    assert document.blocks == ()
    assert document.warnings == ("page:1: no extractable text",)


def test_unsupported_document_type_fails_closed(tmp_path: Path):
    source = tmp_path / "archive.bin"
    source.write_bytes(b"not a supported document")

    with LocalEvidenceIndex(tmp_path / "witness.sqlite3") as index:
        with pytest.raises(UnsupportedDocumentError):
            index_document(source, index)
