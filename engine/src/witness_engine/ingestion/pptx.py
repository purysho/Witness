from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from ..ids import stable_id
from .models import ExtractedBlock, ExtractedDocument


def extract_pptx_file(
    path: str | Path,
    source_version_id: str,
) -> ExtractedDocument:
    """Extract slide text and table rows with stable slide locators."""

    source_path = Path(path)
    presentation = Presentation(str(source_path))
    blocks: list[ExtractedBlock] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        table_number = 0
        for shape_number, shape in enumerate(
            slide.shapes,
            start=1,
        ):
            if getattr(shape, "has_table", False):
                table_number += 1
                for row_number, row in enumerate(
                    shape.table.rows,
                    start=1,
                ):
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]
                    text = " | ".join(cells).strip(" |")
                    if not text:
                        continue
                    locator = (
                        f"slide:{slide_number}/table:{table_number}/"
                        f"row:{row_number}"
                    )
                    blocks.append(
                        ExtractedBlock(
                            block_id=stable_id(
                                "block",
                                source_version_id,
                                "pptx-table-row",
                                locator,
                                text,
                            ),
                            source_version_id=source_version_id,
                            kind="table-row",
                            text=text,
                            locator=locator,
                        )
                    )
                continue

            if not getattr(shape, "has_text_frame", False):
                continue
            text = " ".join(shape.text.split())
            if not text:
                continue
            locator = (
                f"slide:{slide_number}/shape:{shape_number}"
            )
            blocks.append(
                ExtractedBlock(
                    block_id=stable_id(
                        "block",
                        source_version_id,
                        "pptx-text",
                        locator,
                        text,
                    ),
                    source_version_id=source_version_id,
                    kind="slide-text",
                    text=text,
                    locator=locator,
                )
            )

    return ExtractedDocument(
        source_version_id=source_version_id,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "presentationml.presentation"
        ),
        title=source_path.name,
        blocks=tuple(blocks),
    )
