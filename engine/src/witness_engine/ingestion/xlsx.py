from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ..ids import stable_id
from .models import ExtractedBlock, ExtractedDocument


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def extract_xlsx_file(
    path: str | Path,
    source_version_id: str,
) -> ExtractedDocument:
    """Extract worksheet rows with stable sheet/row locators."""

    source_path = Path(path)
    workbook = load_workbook(
        source_path,
        read_only=True,
        data_only=True,
    )
    blocks: list[ExtractedBlock] = []
    try:
        for sheet_index, worksheet in enumerate(
            workbook.worksheets,
            start=1,
        ):
            for row_number, row in enumerate(
                worksheet.iter_rows(values_only=True),
                start=1,
            ):
                cells = [_cell_text(value) for value in row]
                text = " | ".join(cells).strip(" |")
                if not text:
                    continue
                locator = (
                    f"sheet:{sheet_index}/row:{row_number}"
                )
                blocks.append(
                    ExtractedBlock(
                        block_id=stable_id(
                            "block",
                            source_version_id,
                            "xlsx-row",
                            locator,
                            worksheet.title,
                            text,
                        ),
                        source_version_id=source_version_id,
                        kind="table-row",
                        text=text,
                        locator=locator,
                    )
                )
    finally:
        workbook.close()

    return ExtractedDocument(
        source_version_id=source_version_id,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        title=source_path.name,
        blocks=tuple(blocks),
    )
