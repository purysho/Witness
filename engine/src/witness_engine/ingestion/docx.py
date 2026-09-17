from __future__ import annotations

from pathlib import Path

from docx import Document

from ..ids import stable_id
from .models import ExtractedBlock, ExtractedDocument


def extract_docx_file(path: str | Path, source_version_id: str) -> ExtractedDocument:
    """Extract DOCX paragraphs and table rows with stable locators."""

    source_path = Path(path)
    document = Document(str(source_path))
    blocks: list[ExtractedBlock] = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        locator = f"paragraph:{index}"
        style_name = (paragraph.style.name or "").strip().lower() if paragraph.style else ""
        kind = "heading" if style_name.startswith("heading") else "paragraph"
        blocks.append(
            ExtractedBlock(
                block_id=stable_id("block", source_version_id, "docx", locator, text),
                source_version_id=source_version_id,
                kind=kind,
                text=text,
                locator=locator,
            )
        )

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells]
            text = " | ".join(cells).strip(" |")
            if not text:
                continue
            locator = f"table:{table_index}/row:{row_index}"
            blocks.append(
                ExtractedBlock(
                    block_id=stable_id(
                        "block", source_version_id, "docx-table-row", locator, text
                    ),
                    source_version_id=source_version_id,
                    kind="table-row",
                    text=text,
                    locator=locator,
                )
            )

    return ExtractedDocument(
        source_version_id=source_version_id,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        title=source_path.name,
        blocks=tuple(blocks),
    )
