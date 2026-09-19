#!/usr/bin/env python
"""Reproducible mixed-format stress harness for Witness workspaces."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
import sys
import tempfile
import time
import tracemalloc
from collections import Counter
from typing import Callable

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from witness_engine.pipeline import index_document
from witness_engine.rpc.service import RpcService
from witness_engine.tasks import OperationCancelled


FORMATS = ("md", "txt", "html", "csv", "py", "docx", "pptx", "xlsx", "pdf")


def _paragraphs(index: int, count: int) -> list[str]:
    marker = f"STRESS-{index:06d}"
    port = 5200 + (index % 700)
    return [
        (
            f"Witness stress document {index:06d}. "
            f"The marker is {marker}. "
            f"The service port is {port}. "
            f"Paragraph {number:03d} records deterministic corpus evidence."
        )
        for number in range(1, count + 1)
    ]


def _write_markdown(path: Path, paragraphs: list[str]) -> None:
    body = ["# Stress evidence", ""]
    for number, paragraph in enumerate(paragraphs, start=1):
        body.extend([f"## Record {number}", "", paragraph, ""])
    path.write_text("\n".join(body), encoding="utf-8")


def _write_text(path: Path, paragraphs: list[str]) -> None:
    path.write_text("\n\n".join(paragraphs) + "\n", encoding="utf-8")


def _write_html(path: Path, paragraphs: list[str]) -> None:
    payload = "\n".join(f"<p>{paragraph}</p>" for paragraph in paragraphs)
    path.write_text(
        "<!doctype html><html><body><h1>Stress evidence</h1>"
        + payload
        + "</body></html>",
        encoding="utf-8",
    )


def _write_csv(path: Path, paragraphs: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["record", "evidence"])
        for number, paragraph in enumerate(paragraphs, start=1):
            writer.writerow([number, paragraph])


def _write_python(path: Path, paragraphs: list[str]) -> None:
    lines = ["STRESS_RECORDS = ["]
    for paragraph in paragraphs:
        lines.append(f"    {paragraph!r},")
    lines.extend(["]", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_docx(path: Path, paragraphs: list[str]) -> None:
    document = Document()
    document.add_heading("Stress evidence", level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def _write_pptx(path: Path, paragraphs: list[str]) -> None:
    presentation = Presentation()
    for start in range(0, len(paragraphs), 10):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Stress evidence"
        slide.placeholders[1].text = "\n".join(paragraphs[start : start + 10])
    presentation.save(path)


def _write_xlsx(path: Path, paragraphs: list[str]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Stress evidence"
    worksheet.append(["record", "evidence"])
    for number, paragraph in enumerate(paragraphs, start=1):
        worksheet.append([number, paragraph])
    workbook.save(path)
    workbook.close()


def _write_pdf(path: Path, paragraphs: list[str]) -> None:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    for paragraph in paragraphs:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_ref}
                )
            }
        )
        safe = (
            paragraph.replace("\\", "\\\\")
            .replace("(", "\\(")
            .replace(")", "\\)")
        )
        stream = DecodedStreamObject()
        stream.set_data(
            f"BT /F1 11 Tf 72 720 Td ({safe}) Tj ET".encode(
                "latin-1",
                "replace",
            )
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as handle:
        writer.write(handle)


WRITERS: dict[str, Callable[[Path, list[str]], None]] = {
    "md": _write_markdown,
    "txt": _write_text,
    "html": _write_html,
    "csv": _write_csv,
    "py": _write_python,
    "docx": _write_docx,
    "pptx": _write_pptx,
    "xlsx": _write_xlsx,
    "pdf": _write_pdf,
}


def generate_corpus(
    corpus_dir: Path,
    *,
    file_count: int,
    paragraphs_per_file: int,
    large_every: int,
    large_multiplier: int,
) -> tuple[list[Path], Counter[str], int]:
    corpus_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    counts: Counter[str] = Counter()
    large_files = 0
    for index in range(file_count):
        extension = FORMATS[index % len(FORMATS)]
        paragraph_count = paragraphs_per_file
        if large_every > 0 and (index + 1) % large_every == 0:
            paragraph_count *= large_multiplier
            large_files += 1
        path = corpus_dir / f"stress-{index:06d}.{extension}"
        WRITERS[extension](path, _paragraphs(index, paragraph_count))
        paths.append(path)
        counts[extension] += 1
    return paths, counts, large_files


def _database_bytes(workspace: Path) -> int:
    return sum(
        path.stat().st_size
        for path in workspace.glob("witness.db*")
        if path.is_file()
    )


def _counts(service: RpcService) -> tuple[int, int]:
    assert service.lexical is not None
    connection = service.lexical.connection
    sources = int(
        connection.execute(
            "SELECT COUNT(*) FROM source_version_metadata"
        ).fetchone()[0]
    )
    chunks = int(
        connection.execute(
            "SELECT COUNT(*) FROM indexed_chunks"
        ).fetchone()[0]
    )
    return sources, chunks


def _query(service: RpcService) -> dict[str, object]:
    started = time.perf_counter()
    result = service.handle(
        "query.run",
        {"question": "What marker is recorded in stress document 000000?"},
    )
    return {
        "seconds": time.perf_counter() - started,
        "evidence_count": int(result["sufficiency"]["evidence_count"]),
        "citation_count": len(result["answer"]["citations"]),
        "run_id": result["run_id"],
    }


def _repair_check(service: RpcService) -> dict[str, object]:
    healthy = service.handle("workspace.health", {})
    protected_before = healthy["protected_state_fingerprint"]
    assert service.lexical is not None
    with service.lexical.connection:
        service.lexical.connection.execute("DELETE FROM indexed_chunks_fts")
    broken = service.handle("workspace.health", {})
    started = time.perf_counter()
    repaired = service.handle("workspace.repair", {})
    return {
        "status_before": broken["status"],
        "status_after": repaired["after"]["status"],
        "seconds": time.perf_counter() - started,
        "protected_state_unchanged": (
            repaired["protected_state_unchanged"]
            and repaired["after"]["protected_state_fingerprint"]
            == protected_before
        ),
    }


def _cancellation_rollback(root: Path) -> dict[str, object]:
    workspace = root / "cancel-workspace.witness"
    source = root / "cancel-large.md"
    source.write_text(
        "\n\n".join(
            f"Cancellation paragraph {index:04d} keeps deterministic evidence."
            for index in range(600)
        )
        + "\n",
        encoding="utf-8",
    )

    service = RpcService()
    checks = 0
    cancelled = False
    try:
        service.handle("workspace.open", {"path": str(workspace)})
        assert service.lexical is not None
        assert service.vectors is not None
        before = _counts(service)

        def cancel_check() -> None:
            nonlocal checks
            checks += 1
            if checks >= 650:
                raise OperationCancelled("stress-cancel")

        try:
            index_document(
                source,
                service.lexical,
                vector_index=service.vectors,
                embedding_provider=service.embedding_provider,
                visual_index=service.visual_index,
                visual_embedding_provider=service.visual_embedding_provider,
                cancel_check=cancel_check,
            )
        except OperationCancelled:
            cancelled = True

        after = _counts(service)
        health = service.handle("workspace.health", {})
        return {
            "cancelled": cancelled,
            "checks": checks,
            "rolled_back": before == after,
            "health": health["status"],
        }
    finally:
        service.close()


def run(args: argparse.Namespace) -> dict[str, object]:
    if args.files < 1:
        raise ValueError("--files must be >= 1")
    if args.paragraphs < 1:
        raise ValueError("--paragraphs must be >= 1")
    if args.large_every < 0:
        raise ValueError("--large-every must be >= 0")
    if args.large_multiplier < 1:
        raise ValueError("--large-multiplier must be >= 1")
    if args.reopen_cycles < 1:
        raise ValueError("--reopen-cycles must be >= 1")

    temporary: tempfile.TemporaryDirectory[str] | None = None
    if args.root:
        root = Path(args.root).expanduser().resolve()
        if root.exists() and any(root.iterdir()):
            raise ValueError("--root must not contain existing files")
        root.mkdir(parents=True, exist_ok=True)
    elif args.keep:
        root = Path(tempfile.mkdtemp(prefix="witness-stress-")).resolve()
    else:
        temporary = tempfile.TemporaryDirectory(prefix="witness-stress-")
        root = Path(temporary.name).resolve()

    total_started = time.perf_counter()
    try:
        corpus_dir = root / "corpus"
        generation_started = time.perf_counter()
        paths, format_counts, large_files = generate_corpus(
            corpus_dir,
            file_count=args.files,
            paragraphs_per_file=args.paragraphs,
            large_every=args.large_every,
            large_multiplier=args.large_multiplier,
        )
        generation_seconds = time.perf_counter() - generation_started

        workspace = root / "workspace.witness"
        service = RpcService()
        tracemalloc.start()
        import_started = time.perf_counter()
        try:
            service.handle("workspace.open", {"path": str(workspace)})
            for path in paths:
                service.handle("source.import", {"path": str(path)})
            import_seconds = time.perf_counter() - import_started
            _, python_peak_alloc = tracemalloc.get_traced_memory()
            source_versions, chunks = _counts(service)
            health_after_import = service.handle("workspace.health", {})
            query = _query(service)
            repair = _repair_check(service)
        finally:
            tracemalloc.stop()
            service.close()

        cancellation = _cancellation_rollback(root)

        reopen: list[dict[str, object]] = []
        for cycle in range(args.reopen_cycles):
            reopened = RpcService()
            try:
                open_started = time.perf_counter()
                reopened.handle("workspace.open", {"path": str(workspace)})
                open_seconds = time.perf_counter() - open_started
                health = reopened.handle("workspace.health", {})
                query_result = _query(reopened)
                reopen.append(
                    {
                        "cycle": cycle + 1,
                        "open_seconds": open_seconds,
                        "health": health["status"],
                        "query_seconds": query_result["seconds"],
                        "evidence_count": query_result["evidence_count"],
                    }
                )
            finally:
                reopened.close()

        report: dict[str, object] = {
            "schema_version": 1,
            "root": str(root),
            "files_requested": args.files,
            "generated_files": len(paths),
            "format_counts": dict(sorted(format_counts.items())),
            "large_files": large_files,
            "paragraphs_per_file": args.paragraphs,
            "large_every": args.large_every,
            "large_multiplier": args.large_multiplier,
            "generation_seconds": generation_seconds,
            "import_seconds": import_seconds,
            "import_files_per_second": (
                len(paths) / import_seconds if import_seconds > 0 else None
            ),
            "source_versions": source_versions,
            "chunks": chunks,
            "database_bytes": _database_bytes(workspace),
            "python_peak_alloc_bytes": int(python_peak_alloc),
            "health_after_import": health_after_import["status"],
            "query": query,
            "repair": repair,
            "cancellation": cancellation,
            "reopen_cycles": reopen,
            "total_seconds": time.perf_counter() - total_started,
        }

        if args.json_out:
            output = Path(args.json_out).expanduser().resolve()
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(
                json.dumps(report, indent=2, sort_keys=True) + "\n",
                encoding="utf-8",
            )
        print(json.dumps(report, indent=2, sort_keys=True))
        return report
    finally:
        if temporary is not None:
            temporary.cleanup()


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Generate and exercise a reproducible mixed-format Witness corpus."
        )
    )
    parser.add_argument("--files", type=int, default=2000)
    parser.add_argument("--paragraphs", type=int, default=3)
    parser.add_argument("--large-every", type=int, default=50)
    parser.add_argument("--large-multiplier", type=int, default=20)
    parser.add_argument("--reopen-cycles", type=int, default=5)
    parser.add_argument(
        "--root",
        help="Empty directory to retain corpus/workspace artifacts.",
    )
    parser.add_argument(
        "--keep",
        action="store_true",
        help="Keep an automatically-created temporary root.",
    )
    parser.add_argument("--json-out", help="Write the measurement report to JSON.")
    args = parser.parse_args()

    try:
        run(args)
    except Exception as exc:
        print(f"stress harness failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
